from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "생성형AI_활용_프롬프트_정리.md"
OUTPUT = ROOT / "tmp" / "docx_prompt_export" / "FX_OVERDOSE_생성형AI_활용_프롬프트_정리_built.docx"
SKILL_SCRIPTS = Path(
    "/Users/bluem/.codex/plugins/cache/openai-primary-runtime/documents/"
    "26.630.12135/skills/documents/scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry  # noqa: E402


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(85, 85, 85)
PROMPT_INK = RGBColor(32, 33, 36)
EAST_ASIA_FONT = "Apple SD Gothic Neo"


def set_run_font(run, size: float, color: RGBColor = BLACK, bold: bool = False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{key}"), "Arial")
    rfonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def set_style_font(style, size: float, color: RGBColor = BLACK, bold: bool = False) -> None:
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{key}"), "Arial")
    rfonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def set_paragraph_spacing(style, before: float, after: float, line: float) -> None:
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = "DADCE0", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def configure_group_table(table) -> None:
    """Keep a logical prompt block together without drawing an outer frame."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")

    apply_table_geometry(
        table,
        [9360],
        table_width_dxa=9360,
        indent_dxa=0,
        cell_margins_dxa={"top": 0, "bottom": 0, "start": 0, "end": 0},
    )

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_prompt_paragraph_box(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), "F8F9FA")

    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "5")
        el.set(qn("w:color"), "DADCE0")
        borders.append(el)


def add_inline_runs(paragraph, text: str, size: float = 11.0) -> None:
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, max(10.0, size - 0.5), PROMPT_INK)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size)


def next_id(elements, tag: str, attr: str) -> int:
    ids = []
    for element in elements.findall(qn(tag)):
        value = element.get(qn(attr))
        if value is not None and value.lstrip("-").isdigit():
            ids.append(int(value))
    return max(ids, default=0) + 1


def create_numbering_definition(document: Document, bullet: bool) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = next_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = next_id(numbering, "w:num", "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "●" if bullet else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)

    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)

    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{key}"), "Arial")
    rfonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_paragraph(document: Document, text: str, num_id: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    ppr.append(num_pr)
    add_inline_runs(paragraph, text)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    set_style_font(normal, 11)
    set_paragraph_spacing(normal, 0, 8, 1.15)

    heading_tokens = {
        "Heading 1": (20, BLACK, 20, 6),
        "Heading 2": (16, BLACK, 18, 6),
        "Heading 3": (14, RGBColor(67, 67, 67), 16, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        set_style_font(style, size, color, bold=False)
        set_paragraph_spacing(style, before, after, 1.0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    document.core_properties.title = "FX OVERDOSE 생성형 AI 활용 프롬프트 정리서"
    document.core_properties.subject = "제출용 프롬프트 정제본"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""


def add_metadata_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row_data[:2]):
            cell = cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(value)
            set_run_font(run, 10.5, bold=row_index == 0)
    set_table_borders(table)
    apply_table_geometry(
        table,
        [2246, 7114],
        table_width_dxa=9360,
        indent_dxa=0,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def build_document(markdown: str) -> Document:
    document = Document()
    configure_document(document)
    decimal_num_id = create_numbering_definition(document, bullet=False)
    bullet_num_id = create_numbering_definition(document, bullet=True)

    lines = markdown.splitlines()
    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    in_code = False
    group_cell = None
    pending_section_text: str | None = None
    i = 0

    def add_paragraph(style=None):
        target = group_cell if group_cell is not None else document
        return target.add_paragraph(style=style)

    def begin_group(style_name: str, text: str, section_text: str | None = None) -> None:
        nonlocal group_cell
        table = document.add_table(rows=1, cols=1)
        configure_group_table(table)
        group_cell = table.cell(0, 0)
        paragraph = group_cell.paragraphs[0]
        if section_text is not None:
            paragraph.style = document.styles["Heading 1"]
            add_inline_runs(
                paragraph,
                section_text,
                document.styles["Heading 1"].font.size.pt,
            )
            paragraph = group_cell.add_paragraph(style="Heading 2")
        paragraph.style = document.styles[style_name]
        add_inline_runs(paragraph, text, document.styles[style_name].font.size.pt)

    def end_group() -> None:
        nonlocal group_cell
        group_cell = None

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if not paragraph_lines:
            return
        joined = " ".join(line.strip() for line in paragraph_lines)
        paragraph = add_paragraph()
        if joined.startswith("**활용 목적:**") or joined.startswith("아래 내용은 개별 요청"):
            paragraph.paragraph_format.keep_with_next = True
        add_inline_runs(paragraph, joined)
        paragraph_lines = []

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                paragraph = add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.right_indent = Inches(0.18)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(10)
                paragraph.paragraph_format.line_spacing = 1.22
                paragraph.paragraph_format.keep_together = True
                set_prompt_paragraph_box(paragraph)
                for line_index, line in enumerate(code_lines):
                    if line_index:
                        paragraph.add_run().add_break(WD_BREAK.LINE)
                    run = paragraph.add_run(line)
                    set_run_font(run, 10.5, PROMPT_INK)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(raw)
            i += 1
            continue

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            end_group()
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
                add_metadata_table(document, [rows[0], *rows[2:]])
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            end_group()
            level = len(heading.group(1))
            heading_text = heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(heading_text)
                set_run_font(run, 26, BLACK, bold=False)
                subtitle = document.add_paragraph()
                subtitle.paragraph_format.space_before = Pt(0)
                subtitle.paragraph_format.space_after = Pt(18)
                run = subtitle.add_run("제출용 프롬프트 정제본 · 2026년 8월 2일")
                set_run_font(run, 10.5, MUTED)
            else:
                style_name = f"Heading {min(level - 1, 3)}"
                if level == 2 and re.match(r"^(?:[3-9]|10|11)\.\s", heading_text):
                    pending_section_text = heading_text
                    i += 1
                    continue
                should_group = (
                    (level == 2 and heading_text.startswith("2. 공통 프로젝트 컨텍스트"))
                    or (level == 3 and re.match(r"^P-\d{2}\.", heading_text))
                )
                if should_group:
                    begin_group(style_name, heading_text, pending_section_text)
                    pending_section_text = None
                else:
                    paragraph = document.add_paragraph(style=style_name)
                    add_inline_runs(
                        paragraph,
                        heading_text,
                        document.styles[style_name].font.size.pt,
                    )
            i += 1
            continue

        ordered = re.match(r"^\d+\.\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if ordered or bullet:
            flush_paragraph()
            end_group()
            add_list_paragraph(
                document,
                (ordered or bullet).group(1),
                decimal_num_id if ordered else bullet_num_id,
            )
            i += 1
            continue

        paragraph_lines.append(stripped)
        i += 1

    flush_paragraph()
    end_group()
    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    markdown = SOURCE.read_text(encoding="utf-8")
    document = build_document(markdown)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
